"""Tests for reading uploaded CVs into the committee's plain format."""

import io

import pytest
from docx import Document
from fpdf import FPDF

from hcs.parsing import (
    Link,
    body_links,
    identity_header,
    links_from_pdf,
    normalise_cv,
    normalise_dates,
    split_identity,
    text_from_upload,
)
from hcs.profile import build_facts
from hcs.screening import years_of_experience

RAW = """Jane Example
jane@example.com | +44 7700 900123 | linkedin.com/in/jane
Backend Engineer, PayCo
Jan 2021 – Present
• Built the payouts API in Python, used by
  9,000 merchants
• Deploy my services myself
Junior Developer, Agency (2018 - 2021)
- Built Django sites
"""


@pytest.mark.parametrize(
    ("line", "expected"),
    [
        ("Engineer, Shop Jan 2021 – Present", "Engineer, Shop (Jan 2021-present)"),
        ("Engineer, Shop (2018 - 2021)", "Engineer, Shop (2018-2021)"),
        ("Engineer, Shop 03/2019 — current", "Engineer, Shop (Mar 2019-present)"),
        ("Engineer, Shop March 2023 - Feb 2024", "Engineer, Shop (Mar 2023-Feb 2024)"),
        ("Cut costs by 30% in 2023", "Cut costs by 30% in 2023"),
    ],
)
def test_date_ranges_are_written_the_way_the_checks_read_them(line, expected):
    assert normalise_dates(line) == expected


def test_a_pasted_cv_becomes_role_lines_and_bullets_without_identity():
    cv = normalise_cv(RAW)

    assert cv.splitlines() == [
        "Backend Engineer, PayCo (Jan 2021-present)",
        "- Built the payouts API in Python, used by 9,000 merchants",
        "- Deploy my services myself",
        "Junior Developer, Agency (2018-2021)",
        "- Built Django sites",
    ]
    assert years_of_experience(cv, 2026) == 8
    assert [fact.role for fact in build_facts(cv)][
        1
    ] == "Backend Engineer, PayCo (Jan 2021-present)"


def test_a_word_file_is_read():
    document = Document()
    for line in RAW.splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)

    assert normalise_cv(text_from_upload("cv.docx", buffer.getvalue())) == normalise_cv(RAW)


def minimal_pdf(lines: list[str]) -> bytes:
    """A one-page PDF with the lines as text, written by hand (no PDF writer needed)."""
    text = " ".join(f"({line}) Tj T*" for line in lines)
    stream = f"BT /F1 11 Tf 14 TL 50 750 Td {text} ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)
    return out


def test_a_pdf_is_read():
    pdf = minimal_pdf(["Backend Engineer, PayCo 2021 - Present", "- Built the payouts API"])

    assert normalise_cv(text_from_upload("cv.pdf", pdf)).splitlines() == [
        "Backend Engineer, PayCo (2021-present)",
        "- Built the payouts API",
    ]


def test_runs_of_spaces_from_a_pdf_layout_become_one():
    assert normalise_cv("Backend Engineer, PayCo    2021 - Present") == (
        "Backend Engineer, PayCo (2021-present)"
    )


def test_the_removed_identity_lines_are_kept_for_the_pdf_header():
    cv, removed = split_identity(RAW)

    assert cv == normalise_cv(RAW)
    assert removed == ["Jane Example", "jane@example.com | +44 7700 900123 | linkedin.com/in/jane"]


def test_the_header_puts_the_name_first_and_one_contact_item_per_line():
    removed = ["jane@example.com · +44 7700 900123", "Jane Example", "github.com/jane"]

    assert identity_header(removed).splitlines() == [
        "Jane Example",
        "jane@example.com",
        "+44 7700 900123",
        "github.com/jane",
    ]


def test_a_line_the_pdf_wrapped_is_joined_back_but_short_lines_are_not():
    raw = "\n".join(
        [
            "Python backend engineer with hands-on experience building production software, improving",
            "database performance and security.",
            "Backend Engineer, PayCo (2021-present)",
            "- Cut search time 5x over a scan layer that spans the databases we support, including",
            "PostgreSQL, MySQL and Oracle.",
            "- Built the payouts API",
            "Skills",
            "Languages:Python, SQL",
        ]
    )

    assert normalise_cv(raw).splitlines() == [
        "Python backend engineer with hands-on experience building production software, improving"
        " database performance and security.",
        "Backend Engineer, PayCo (2021-present)",
        "- Cut search time 5x over a scan layer that spans the databases we support, including"
        " PostgreSQL, MySQL and Oracle.",
        "- Built the payouts API",
        "Skills",
        "Languages: Python, SQL",
    ]


def test_a_degree_wrapped_over_two_lines_becomes_one_dated_line():
    raw = "\n".join(
        [
            "Education",
            "B.Sc. Computer Science (AI specialisation), Faculty of Computers and Artificial",
            "Intelligence, Example University (2021-2025)",
        ]
    )

    assert normalise_cv(raw).splitlines()[1] == (
        "B.Sc. Computer Science (AI specialisation), Faculty of Computers and Artificial"
        " Intelligence, Example University (2021-2025)"
    )


def test_only_the_name_is_removed_not_the_headline_or_an_about_me_heading():
    raw = "\n".join(
        [
            "Jane Example",
            "Software Engineer",
            "Lisbon |jane@example.com | LinkedIn | GitHub",
            "About Me",
            "Backend engineer who builds compliance software.",
            "Experience",
            "Software Engineer, Northwind (2026-present)",
        ]
    )

    cv, removed = split_identity(raw)

    assert removed == ["Jane Example", "Lisbon |jane@example.com | LinkedIn | GitHub"]
    assert cv.splitlines()[:3] == ["Software Engineer", "About Me", cv.splitlines()[2]]


def pdf_with_links() -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=11)
    pdf.cell(40, 8, "LinkedIn", link="https://www.linkedin.com/in/jane/")
    pdf.cell(40, 8, "GitHub", link="https://github.com/jane", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 8, "Doctors Hub /github", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(30, 8, "/github", link="https://github.com/jane/doctors-hub")
    return bytes(pdf.output())


def test_links_are_read_from_a_pdf_with_the_text_they_sit_on():
    links = links_from_pdf(pdf_with_links())

    assert [link.uri for link in links] == [
        "https://www.linkedin.com/in/jane/",
        "https://github.com/jane",
        "https://github.com/jane/doctors-hub",
    ]
    assert links[0].text == "LinkedIn"


def test_a_contact_item_that_was_a_link_becomes_its_address():
    links = [
        Link("mailto:jane@example.com", "jane@example.com"),
        Link("https://www.linkedin.com/in/jane/", "LinkedIn"),
        Link("https://github.com/jane", "GitHub"),
    ]

    header = identity_header(["Jane Example", "jane@example.com | LinkedIn | GitHub"], links)

    assert header.splitlines() == [
        "Jane Example",
        "jane@example.com",
        "linkedin.com/in/jane",
        "github.com/jane",
    ]


def test_repository_links_go_to_the_cv_labels_in_order_and_contact_links_stay_out():
    cv = "Video QA /github\n- Stack: FastAPI\nDoctors Hub /github\n- Stack: Django"
    links = [
        Link("https://github.com/jane", "GitHub"),
        Link("https://github.com/jane/video-qa", ""),
        Link("https://github.com/jane/doctors-hub", ""),
    ]

    assert body_links(cv, links, header="github.com/jane") == [
        ("/github", "https://github.com/jane/video-qa"),
        ("/github", "https://github.com/jane/doctors-hub"),
    ]
