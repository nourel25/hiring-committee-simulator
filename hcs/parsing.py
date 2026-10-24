"""Reading an uploaded CV (PDF, Word or text) into the plain format the committee reads: a dated role
line such as "Backend Engineer, shop (2020-present)", then one bullet per line.

Extraction is never trusted blindly: the UI shows the result for the user to correct before any
model reads it. Scanned PDFs (images, no text layer) and two-column layouts come out badly.
"""

import io
import re
from dataclasses import dataclass

from hcs.screening import MONTHS, PERIOD, is_heading

# "Jan 2021 – Present", "2019 - 2021", "03/2018 — current": a date range in any common form.
# Each side captures a month name or a month number, then the year.
MONTH = r"(?:(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+|(\d{1,2})/)?"
DATE_RANGE = re.compile(
    rf"\(?{MONTH}((?:19|20)\d{{2}})\s*(?:-|–|—|to)\s*{MONTH}"
    r"((?:19|20)\d{2}|present|current|now|today)\s*\)?",
    re.IGNORECASE,
)
BULLET = re.compile(r"^\s*(?:[•●▪◦‣∙·*]|-(?=\s)|–(?=\s)|\d+[.)](?=\s))\s*")
EMAIL = re.compile(r"\S+@\S+\.\w+")
PHONE = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")
LINK = re.compile(r"(?:https?://|www\.|linkedin\.com|github\.com)\S*", re.IGNORECASE)
# A PDF wraps a line near the page width: about 85-120 characters in the CVs seen so far.
WRAP_WIDTH = 70
NAME_LINE = re.compile(r"^[A-Z][a-zA-Z'.-]+(?:\s+[A-Z][a-zA-Z'.-]+){1,3}$")


def text_from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def text_from_docx(data: bytes) -> str:
    from docx import Document

    return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)


@dataclass
class Link:
    """A hyperlink in the uploaded file: its target and, when it can be read, the text it sits on."""

    uri: str
    text: str = ""


def links_from_pdf(data: bytes) -> list[Link]:
    """Every URI link annotation, in page order, with the text drawn inside its rectangle. PDFs
    keep links apart from the text, so text extraction alone loses them."""
    from pypdf import PdfReader

    found = []
    for page in PdfReader(io.BytesIO(data)).pages:
        fragments: list[tuple[float, float, str]] = []

        def collect(text, cm, tm, font, size, fragments=fragments):
            if text.strip():
                x = tm[4] * cm[0] + tm[5] * cm[2] + cm[4]
                y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
                fragments.append((x, y, text.strip()))

        page.extract_text(visitor_text=collect)
        annotations = [annotation.get_object() for annotation in page.get("/Annots") or []]
        for annotation in annotations:
            action = annotation.get("/A")
            if not action or "/URI" not in action:
                continue
            x0, y0, x1, y1 = (float(value) for value in annotation["/Rect"])
            inside = [
                text for x, y, text in fragments if x0 - 2 <= x <= x1 + 2 and y0 - 3 <= y <= y1 + 3
            ]
            found.append(Link(str(action["/URI"]), " ".join(inside)))
    return found


def links_from_docx(data: bytes) -> list[Link]:
    from docx import Document

    return [
        Link(link.address, link.text)
        for paragraph in Document(io.BytesIO(data)).paragraphs
        for link in paragraph.hyperlinks
        if link.address
    ]


def links_from_upload(name: str, data: bytes) -> list[Link]:
    suffix = name.rsplit(".", 1)[-1].casefold()
    if suffix == "pdf":
        return links_from_pdf(data)
    if suffix == "docx":
        return links_from_docx(data)
    return []


def short_url(uri: str) -> str:
    """'https://www.linkedin.com/in/jane/' -> 'linkedin.com/in/jane': what a CV shows."""
    return re.sub(r"^(?:https?://)?(?:www\.)?", "", uri).rstrip("/")


def text_from_upload(name: str, data: bytes) -> str:
    """The raw text of an uploaded file, by its extension."""
    suffix = name.rsplit(".", 1)[-1].casefold()
    if suffix == "pdf":
        return text_from_pdf(data)
    if suffix == "docx":
        return text_from_docx(data)
    return data.decode("utf-8", errors="replace")


def normalise_dates(line: str) -> str:
    """Write any date range as "(2021-present)", or "(Mar 2021-present)" when it has months: the
    form the year and role checks read. Months are kept, so 10 months don't count as a year."""

    def month(name: str | None, number: str | None) -> str:
        if name:
            return name[:3].title() + " "
        if number and 1 <= int(number) <= 12:
            return MONTHS[int(number) - 1].title() + " "
        return ""

    def replace(match: re.Match) -> str:
        start_name, start_number, start, end_name, end_number, end = match.groups()
        end = end.casefold()
        if end in {"present", "current", "now", "today"}:
            return f"({month(start_name, start_number)}{start}-present)"
        return f"({month(start_name, start_number)}{start}-{month(end_name, end_number)}{end})"

    return DATE_RANGE.sub(replace, line)


def joins_previous(line: str, previous: str) -> bool:
    """A line the PDF's layout wrapped off the end of the previous one. Wrapping happens near the
    full width, so the previous line is long and doesn't end a sentence; headings, dated role lines
    and "Label: values" lines are never continued."""
    if not previous or previous.endswith((".", ":", ";", "!", "?")) or len(previous) < WRAP_WIDTH:
        return False
    if is_heading(previous) or PERIOD.search(previous) or BULLET.match(line) or is_heading(line):
        return False
    return not re.match(r"[A-Z][\w &/+-]{0,30}:", line)


def looks_like_phone(line: str) -> bool:
    """A run of 9 or more digits with phone punctuation, on a short line ("+49 151 2345 6789")."""
    match = PHONE.search(line)
    return bool(match) and sum(c.isdigit() for c in match.group()) >= 9 and len(line.split()) <= 6


def is_identity(line: str, before_first_role: bool) -> bool:
    """Contact details, links, and a bare name above the first role: removed for fair review.
    A section heading ("About Me") is never a name."""
    if EMAIL.search(line) or LINK.search(line) or looks_like_phone(line):
        return True
    return before_first_role and bool(NAME_LINE.match(line.strip())) and not is_heading(line)


def normalise_cv(text: str) -> str:
    """Bullets as "- ", date ranges as "(YYYY-YYYY)", wrapped lines joined, identity lines removed."""
    return split_identity(text)[0]


def split_identity(text: str) -> tuple[str, list[str]]:
    """The normalised CV, and the identity lines removed from it, in order, so the user can put
    them back on the exported PDF (they never reach a model)."""
    lines: list[str] = []
    removed: list[str] = []
    seen_role = False
    name_removed = False
    for raw in text.splitlines():
        line = normalise_dates(re.sub(r"\s{2,}", " ", raw.strip()))
        # "Languages:Python" -> "Languages: Python" (not "10:30", not a link: links are removed).
        line = re.sub(r"(?<=[A-Za-z]):(?=[A-Za-z])", ": ", line)
        if not line:
            if lines and lines[-1]:
                lines.append("")
            continue
        is_role = bool(PERIOD.search(line))
        if PERIOD.fullmatch(line) and lines and lines[-1] and not lines[-1].startswith("- "):
            # Dates on their own line under the title: put them on the title's line.
            lines[-1] += " " + line
            seen_role = True
            continue
        # Only one bare name is removed, and only before any kept line: a headline under the name
        # ("Software Engineer") looks like a name too, and the recruiter's title check reads it.
        before_first_role = not seen_role and not lines and not name_removed
        if not is_role and is_identity(line, before_first_role):
            name_removed = name_removed or bool(NAME_LINE.match(line))
            removed.append(line)
            continue
        seen_role = seen_role or is_role
        bullet = BULLET.match(line)
        if bullet:
            lines.append("- " + line[bullet.end() :])
        elif lines and lines[-1].startswith("- ") and (line[0].islower() or line[0].isdigit()):
            # A bullet wrapped onto the next line by the PDF's layout.
            lines[-1] += " " + line
        elif lines and joins_previous(line, lines[-1]):
            lines[-1] += " " + line
        else:
            lines.append(line)
    return "\n".join(lines).strip(), removed


def identity_header(removed: list[str], links: list[Link] = ()) -> str:
    """The removed lines in the PDF header's format: the name first, then one contact item per
    line ("jane@example.com | +44 7700 900123" becomes two lines). An item that was only the text
    of a link ("LinkedIn", "GitHub") becomes the link's address, so it works on the new PDF."""
    items = [item.strip() for line in removed for item in re.split(r"[|•·]", line) if item.strip()]
    for index, item in enumerate(items):
        word = item.casefold().replace(" ", "")
        match = next(
            (
                link
                for link in links
                if not link.uri.startswith(("mailto:", "tel:"))
                and (link.text.casefold().replace(" ", "") == word or word in link.uri.casefold())
            ),
            None,
        )
        if match and not LINK.search(item):
            items[index] = short_url(match.uri)
    names = [item for item in items if NAME_LINE.match(item)][:1]
    return "\n".join(names + [item for item in items if item not in names])


def body_links(cv: str, links: list[Link], header: str) -> list[tuple[str, str]]:
    """Links that belong in the CV body, as (text it sits on, address), in CV order: a project's
    "/github" or a certificate line. Contact links are the header's; a repository link whose text
    couldn't be read is paired, in order, with the CV's "/github" or "GitHub" labels."""
    header_text = header.casefold()
    placed: list[tuple[str, str]] = []
    labels = [
        match.group()
        for match in re.finditer(r"/github\b|\bGitHub\b(?! Actions)", cv, re.IGNORECASE)
    ]
    for link in links:
        if (
            link.uri.startswith(("mailto:", "tel:"))
            or short_url(link.uri).casefold() in header_text
        ):
            continue
        text = link.text.strip()
        if len(text) >= 3 and text in cv:
            placed.append((text, link.uri))
        elif re.search(r"github\.com/[^/]+/[^/]+", link.uri) and labels:
            placed.append((labels.pop(0), link.uri))
    return placed
